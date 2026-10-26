#!/usr/bin/env python3
"""
Génère data/products.js depuis catalogue.xlsx.

Usage :
    python3 tools/excel_to_products.py catalogue.xlsx data/products.js

Le fichier Excel est la source de vérité. Ne modifie pas data/products.js à la main :
il sera régénéré à chaque déploiement Netlify.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent.parent

SHOE_SIZES = ["24", "26", "28", "30", "32", "34", "36", "37", "38", "39", "40", "41", "42", "43", "44", "45"]
BELT_SIZES = ["75", "80", "85", "90", "95", "100", "105"]
STOCK_KEYS = [f"stock_{s}" for s in SHOE_SIZES] + [f"stock_{s}cm" for s in BELT_SIZES] + ["stock_unique"]

VALID_CATEGORIES = ["babouches", "ceintures", "sacs", "poufs", "accessoires"]
VALID_AUDIENCES = ["femme", "homme", "enfant"]
PLACEHOLDER = "assets/img/products/placeholder.svg"
DEFAULT_HEX = "#6B4423"

CATEGORIES = [
    {"id": "all", "label": {"fr": "Tout", "en": "All"}},
    {"id": "babouches", "label": {"fr": "Babouches", "en": "Babouches"}},
    {"id": "babouches-femme", "label": {"fr": "Babouches femme", "en": "Women's babouches"}},
    {"id": "babouches-homme", "label": {"fr": "Babouches homme", "en": "Men's babouches"}},
    {"id": "babouches-enfant", "label": {"fr": "Babouches enfant", "en": "Kids' babouches"}},
    {"id": "ceintures", "label": {"fr": "Ceintures", "en": "Belts"}},
    {"id": "sacs", "label": {"fr": "Sacs", "en": "Bags"}},
    {"id": "poufs", "label": {"fr": "Poufs", "en": "Poufs"}},
    {"id": "accessoires", "label": {"fr": "Accessoires", "en": "Accessories"}},
]


def text(value) -> str:
    return "" if value is None else str(value).strip()


def truthy(value) -> bool:
    return text(value).upper() in {"OUI", "YES", "TRUE", "VRAI", "1"}


def as_int(value, label: str, errors: list[str], row: int) -> int:
    raw = text(value).replace(" ", "").replace(",", ".")
    if raw == "":
        errors.append(f"Ligne {row} : {label} est vide.")
        return 0
    try:
        number = int(float(raw))
    except (TypeError, ValueError):
        errors.append(f"Ligne {row} : {label} doit être un nombre entier (valeur reçue : {value!r}).")
        return 0
    if number < 0:
        errors.append(f"Ligne {row} : {label} ne peut pas être négatif.")
        return 0
    return number


def as_price(value, errors: list[str], row: int):
    raw = text(value).replace(" ", "").replace(",", ".")
    if raw == "":
        errors.append(f"Ligne {row} : le prix est obligatoire.")
        return 0
    try:
        number = float(raw)
    except ValueError:
        errors.append(f"Ligne {row} : prix invalide ({value!r}).")
        return 0
    if number < 0:
        errors.append(f"Ligne {row} : le prix ne peut pas être négatif.")
        return 0
    return int(number) if number.is_integer() else round(number, 2)


def image_path(value, errors: list[str], row: int, warnings: list[str]) -> str:
    raw = text(value)
    if not raw:
        warnings.append(f"Ligne {row} : image_1 vide, image « Photo à venir » utilisée.")
        return PLACEHOLDER
    if raw.startswith(("http://", "https://")):
        return raw
    clean = raw[2:] if raw.startswith("./") else raw
    if not (REPO / clean).is_file():
        warnings.append(
            f"Ligne {row} : image introuvable « {clean} » — image « Photo à venir » utilisée. "
            "Vérifie que la photo a bien été ajoutée à assets/img/products/."
        )
        return PLACEHOLDER
    return clean


def main() -> int:
    if len(sys.argv) < 3:
        print("Usage : python3 tools/excel_to_products.py catalogue.xlsx data/products.js", file=sys.stderr)
        return 1

    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    if not source.is_absolute():
        source = (Path.cwd() / source).resolve()
    if not target.is_absolute():
        target = (Path.cwd() / target).resolve()

    if not source.is_file():
        print(f"ERREUR : fichier Excel introuvable : {source}", file=sys.stderr)
        return 1

    workbook = load_workbook(source, data_only=True)
    if "Produits" in workbook.sheetnames:
        sheet = workbook["Produits"]
    else:
        sheet = workbook.worksheets[0]

    headers: dict[str, int] = {}
    for cell in sheet[1]:
        key = text(cell.value)
        if key:
            headers[key] = cell.column

    required = ["id", "nom_fr", "prix_mad", "categorie"]
    missing = [key for key in required if key not in headers]
    if missing:
        print(
            "ERREUR : colonnes obligatoires manquantes dans l'onglet Produits : " + ", ".join(missing),
            file=sys.stderr,
        )
        return 1

    def value(row: int, key: str):
        column = headers.get(key)
        return None if column is None else sheet.cell(row=row, column=column).value

    products = []
    errors: list[str] = []
    warnings: list[str] = []
    seen_ids: set[str] = set()

    for row in range(2, sheet.max_row + 1):
        product_id = text(value(row, "id"))
        name_fr = text(value(row, "nom_fr"))
        if not product_id and not name_fr:
            continue

        if not product_id:
            errors.append(f"Ligne {row} : l'ID est obligatoire.")
            continue
        if not re.fullmatch(r"[a-z0-9][a-z0-9-]*", product_id):
            errors.append(
                f"Ligne {row} : ID « {product_id} » invalide. Utilise des minuscules, chiffres et tirets "
                "(ex. bab-013)."
            )
            continue
        if product_id in seen_ids:
            errors.append(f"Ligne {row} : l'ID « {product_id} » est déjà utilisé.")
            continue
        seen_ids.add(product_id)

        if not name_fr:
            errors.append(f"Ligne {row} : le nom en français est obligatoire.")
            continue

        name_en = text(value(row, "nom_en")) or name_fr
        desc_fr = text(value(row, "description_fr"))
        desc_en = text(value(row, "description_en")) or desc_fr

        category = text(value(row, "categorie")).lower()
        if category not in VALID_CATEGORIES:
            errors.append(
                f"Ligne {row} : catégorie « {category} » invalide. Valeurs autorisées : "
                + ", ".join(VALID_CATEGORIES)
            )
            continue

        audience = text(value(row, "audience")).lower()
        if category == "babouches":
            if audience and audience not in VALID_AUDIENCES:
                errors.append(
                    f"Ligne {row} : public « {audience} » invalide. Valeurs autorisées : "
                    + ", ".join(VALID_AUDIENCES)
                )
                continue
            if not audience:
                warnings.append(
                    f"Ligne {row} : public vide pour une babouche — elle n'apparaîtra pas dans "
                    "Babouches femme/homme/enfant."
                )
        else:
            audience = ""

        price = as_price(value(row, "prix_mad"), errors, row)

        color_fr = text(value(row, "couleur_fr")) or "—"
        color_en = text(value(row, "couleur_en")) or color_fr
        color_hex = text(value(row, "couleur_hex")) or DEFAULT_HEX
        if not re.fullmatch(r"#[0-9A-Fa-f]{6}", color_hex):
            warnings.append(f"Ligne {row} : couleur hex « {color_hex} » invalide — {DEFAULT_HEX} utilisée.")
            color_hex = DEFAULT_HEX

        image_1 = image_path(value(row, "image_1"), errors, row, warnings)
        image_2_raw = text(value(row, "image_2"))
        images = [image_1]
        if image_2_raw:
            image_2 = image_path(image_2_raw, errors, row, warnings)
            if image_2 not in images:
                images.append(image_2)

        one_size = truthy(value(row, "taille_unique"))
        featured = truthy(value(row, "mise_en_avant"))

        sizes = []
        if one_size:
            stock = as_int(value(row, "stock_unique"), "stock_unique", errors, row)
            sizes = [{"size": "one", "stock": stock}]
        else:
            for key in STOCK_KEYS:
                if key == "stock_unique":
                    continue
                cell_value = value(row, key)
                if text(cell_value) == "":
                    continue
                stock = as_int(cell_value, key, errors, row)
                if key.endswith("cm"):
                    label = key[len("stock_"):-2] + " cm"
                else:
                    label = key[len("stock_"):]
                sizes.append({"size": label, "stock": stock})
            if not sizes:
                errors.append(
                    f"Ligne {row} : aucune taille/stock renseigné. Remplis au moins une colonne Stock "
                    "(ou coche Taille unique)."
                )
                continue

        product = {
            "id": product_id,
            "name": {"fr": name_fr, "en": name_en},
            "description": {"fr": desc_fr, "en": desc_en},
            "price": price,
            "category": category,
        }
        if audience:
            product["audience"] = audience
        product["images"] = images
        product["color"] = {"fr": color_fr, "en": color_en, "hex": color_hex}
        product["sizes"] = sizes
        if one_size:
            product["oneSize"] = True
        product["featured"] = featured
        products.append(product)

    if errors:
        print("ERREUR : le catalogue contient des erreurs, le fichier n'a pas été généré.", file=sys.stderr)
        for error in errors:
            print("  - " + error, file=sys.stderr)
        return 1

    if not products:
        print("ERREUR : aucun produit trouvé dans catalogue.xlsx.", file=sys.stderr)
        return 1

    header = """/* ============================================================
   CATALOGUE GÉNÉRÉ AUTOMATIQUEMENT DEPUIS catalogue.xlsx
   Ne pas modifier ce fichier à la main.
   Modifie catalogue.xlsx puis lance :
     python3 tools/excel_to_products.py catalogue.xlsx data/products.js
   ============================================================ */

"""
    content = header
    content += "const CATEGORIES = " + json.dumps(CATEGORIES, ensure_ascii=False, indent=2) + ";\n\n"
    content += "const PRODUCTS = " + json.dumps(products, ensure_ascii=False, indent=2) + ";\n"

    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(content, encoding="utf-8")

    counts: dict[str, int] = {}
    for product in products:
        counts[product["category"]] = counts.get(product["category"], 0) + 1

    print(f"{len(products)} produits générés vers {target}")
    print("  " + " · ".join(f"{k}: {v}" for k, v in sorted(counts.items())))
    for warning in warnings:
        print("  ATTENTION : " + warning, file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
